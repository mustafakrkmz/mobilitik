import csv
import sqlite3

from openpyxl import load_workbook

from mobilitik.desktop.data import ComplaintRepository


def _seed(db_path):
    with sqlite3.connect(db_path) as conn:
        conn.executemany(
            """
            INSERT INTO complaints (
                company, complaint_url, complaint_date, title, complaint_text,
                resolved, company_responded, company_response_text,
                company_response_date, response_hours,
                resolution_text, resolution_date, resolution_hours,
                listing_page, scraped_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    "istikbal",
                    "https://example.test/1",
                    "2026-03-01 10:00:00",
                    "Teslimat gecikti",
                    "Teslimat tarihi ertelendi",
                    1,
                    1,
                    "İlgileniyoruz",
                    "2026-03-01 16:00:00",
                    6.0,
                    "Sorun çözüldü",
                    "2026-03-03 10:00:00",
                    48.0,
                    1,
                    "2026-03-01T11:00:00",
                ),
                (
                    "bellona",
                    "https://example.test/2",
                    "2026-04-01 10:00:00",
                    "Koltuk kumaşı yırtıldı",
                    "Kumaş deforme oldu",
                    0,
                    0,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    1,
                    "2026-04-01T11:00:00",
                ),
            ],
        )


def test_repository_filters_and_metrics(tmp_path):
    db_path = tmp_path / "mobilitik.db"
    repo = ComplaintRepository(db_path)
    _seed(db_path)

    metrics = repo.metrics()
    assert metrics == {
        "total": 2,
        "resolved": 1,
        "responded": 1,
        "timed_responses": 1,
        "timed_resolutions": 1,
        "median_response_hours": 6.0,
        "median_resolution_hours": 48.0,
    }

    rows = repo.filtered_complaints(
        company="istikbal", start_date="2026-01-01", end_date="2026-03-31"
    )
    assert len(rows) == 1
    assert rows[0]["title"] == "Teslimat gecikti"

    filtered_metrics = repo.filtered_metrics(
        company="istikbal", start_date="2026-01-01", end_date="2026-03-31"
    )
    assert filtered_metrics["median_response_hours"] == 6.0
    assert filtered_metrics["median_resolution_hours"] == 48.0


def test_repository_csv_and_excel_export(tmp_path):
    db_path = tmp_path / "mobilitik.db"
    repo = ComplaintRepository(db_path)
    _seed(db_path)

    csv_path = tmp_path / "out.csv"
    xlsx_path = tmp_path / "out.xlsx"
    repo.export_csv(csv_path)
    repo.export_xlsx(xlsx_path)

    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 2
    assert {row["company"] for row in rows} == {"istikbal", "bellona"}
    istikbal = next(row for row in rows if row["company"] == "istikbal")
    assert istikbal["response_hours"] == "6.0"
    assert istikbal["resolution_hours"] == "48.0"

    workbook = load_workbook(xlsx_path, read_only=True)
    sheet = workbook["Şikâyetler"]
    assert sheet.max_row == 3
    assert sheet["B2"].value in {"istikbal", "bellona"}


def test_repository_filtered_csv_and_excel_export(tmp_path):
    db_path = tmp_path / "mobilitik.db"
    repo = ComplaintRepository(db_path)
    _seed(db_path)

    csv_path = tmp_path / "filtered.csv"
    xlsx_path = tmp_path / "filtered.xlsx"

    # Sadece istikbal filtresiyle export
    repo.export_csv(csv_path, company="istikbal")
    repo.export_xlsx(xlsx_path, company="istikbal")

    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 1
    assert rows[0]["company"] == "istikbal"
    assert rows[0]["title"] == "Teslimat gecikti"

    workbook = load_workbook(xlsx_path, read_only=True)
    sheet = workbook["Şikâyetler"]
    assert sheet.max_row == 2
    assert sheet["B2"].value == "istikbal"


def test_repository_default_db_path():
    from mobilitik.config import DEFAULT_DB_PATH

    repo = ComplaintRepository()
    assert repo.db_path == DEFAULT_DB_PATH


def test_repository_indexes_created(tmp_path):
    db_path = tmp_path / "mobilitik.db"
    ComplaintRepository(db_path)
    with sqlite3.connect(db_path) as conn:
        indexes = {
            row[1]
            for row in conn.execute("SELECT type, name FROM sqlite_master WHERE type='index'").fetchall()
        }
    assert "idx_complaints_company_date" in indexes
    assert "idx_complaints_resolved" in indexes
