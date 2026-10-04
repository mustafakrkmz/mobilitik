from __future__ import annotations

import csv
import sqlite3
from pathlib import Path

from mobilitik.company import normalize_company_input
from mobilitik.text_quality import normalize_space, sanitize_complaint_body
from mobilitik.timing import median


COLUMNS = [
    "id",
    "company",
    "complaint_url",
    "complaint_date",
    "title",
    "complaint_text",
    "resolved",
    "company_responded",
    "company_response_text",
    "company_response_date",
    "response_hours",
    "resolution_text",
    "resolution_date",
    "resolution_hours",
    "listing_page",
    "scraped_at",
]

TIMING_COLUMNS = {
    "company_response_date": "TEXT",
    "response_hours": "REAL",
    "resolution_text": "TEXT",
    "resolution_date": "TEXT",
    "resolution_hours": "REAL",
}


class ComplaintRepository:
    def __init__(self, db_path: Path | None = None):
        if db_path is None:
            from mobilitik.config import DEFAULT_DB_PATH
            db_path = DEFAULT_DB_PATH
        self.db_path = Path(db_path)
        self._ensure_database()

    def _connect(self):
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        return conn

    @staticmethod
    def _table_exists(conn: sqlite3.Connection, table: str) -> bool:
        return conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
            (table,),
        ).fetchone() is not None

    def _ensure_database(self):
        with self._connect() as conn:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS complaints (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    company TEXT NOT NULL,
                    complaint_url TEXT NOT NULL UNIQUE,
                    complaint_date TEXT,
                    title TEXT,
                    complaint_text TEXT,
                    resolved INTEGER NOT NULL DEFAULT 0,
                    company_responded INTEGER NOT NULL DEFAULT 0,
                    company_response_text TEXT,
                    company_response_date TEXT,
                    response_hours REAL,
                    resolution_text TEXT,
                    resolution_date TEXT,
                    resolution_hours REAL,
                    listing_page INTEGER,
                    scraped_at TEXT NOT NULL
                )
                """
            )
            existing = {
                row[1]
                for row in conn.execute("PRAGMA table_info(complaints)").fetchall()
            }
            for column, sql_type in TIMING_COLUMNS.items():
                if column not in existing:
                    conn.execute(f"ALTER TABLE complaints ADD COLUMN {column} {sql_type}")

            conn.execute(
                "CREATE INDEX IF NOT EXISTS idx_complaints_company_date ON complaints(company, complaint_date)"
            )
            conn.execute(
                "CREATE INDEX IF NOT EXISTS idx_complaints_resolved ON complaints(resolved)"
            )

            self._repair_legacy_text_rows(conn)

    def _repair_legacy_text_rows(self, conn: sqlite3.Connection):
        """Remove SEO copy and company answers that leaked into old complaint bodies."""
        rows = conn.execute(
            "SELECT id, complaint_text, company_response_text FROM complaints "
            "WHERE complaint_text IS NOT NULL"
        ).fetchall()
        changed_ids: list[int] = []
        for row in rows:
            original = normalize_space(row["complaint_text"]) or None
            cleaned = sanitize_complaint_body(
                row["complaint_text"],
                row["company_response_text"],
            ) or None
            if cleaned != original:
                conn.execute(
                    "UPDATE complaints SET complaint_text = ? WHERE id = ?",
                    (cleaned, row["id"]),
                )
                changed_ids.append(int(row["id"]))

        if not changed_ids:
            return

        placeholders = ",".join("?" for _ in changed_ids)
        for table in ("sentiment_results", "aspect_sentiment"):
            if self._table_exists(conn, table):
                conn.execute(
                    f"DELETE FROM {table} WHERE complaint_id IN ({placeholders})",
                    changed_ids,
                )

    @staticmethod
    def _metrics_from_rows(rows):
        rows = list(rows)
        total = len(rows)
        return {
            "total": total,
            "resolved": sum(int(bool(row["resolved"])) for row in rows),
            "responded": sum(int(bool(row["company_responded"])) for row in rows),
            "timed_responses": sum(row["response_hours"] is not None for row in rows),
            "timed_resolutions": sum(row["resolution_hours"] is not None for row in rows),
            "median_response_hours": median(row["response_hours"] for row in rows),
            "median_resolution_hours": median(row["resolution_hours"] for row in rows),
        }

    def metrics(self):
        return self._metrics_from_rows(self.all_complaints())

    def filtered_metrics(self, company: str | None = None, start_date: str | None = None, end_date: str | None = None):
        return self._metrics_from_rows(
            self.filtered_complaints(company=company, start_date=start_date, end_date=end_date)
        )

    def list_complaints(self, limit: int = 500):
        with self._connect() as conn:
            rows = conn.execute(
                """
                SELECT *
                FROM complaints
                ORDER BY COALESCE(complaint_date, scraped_at) DESC, id DESC
                LIMIT ?
                """,
                (limit,),
            ).fetchall()
        return rows

    def filtered_complaints(self, company: str | None = None, start_date: str | None = None, end_date: str | None = None):
        clauses = []
        params: list[str] = []

        if company:
            clauses.append("company = ?")
            params.append(normalize_company_input(company))
        if start_date:
            clauses.append("date(complaint_date) >= date(?)")
            params.append(start_date)
        if end_date:
            clauses.append("date(complaint_date) <= date(?)")
            params.append(end_date)

        where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
        query = f"""
            SELECT *
            FROM complaints
            {where}
            ORDER BY COALESCE(complaint_date, scraped_at) DESC, id DESC
        """
        with self._connect() as conn:
            return conn.execute(query, params).fetchall()

    def all_complaints(self):
        return self.filtered_complaints()

    def export_csv(
        self,
        path: Path,
        *,
        company: str | None = None,
        start_date: str | None = None,
        end_date: str | None = None,
        records: list | None = None,
    ):
        if records is not None:
            rows = records
        elif company or start_date or end_date:
            rows = self.filtered_complaints(company=company, start_date=start_date, end_date=end_date)
        else:
            rows = self.all_complaints()

        with Path(path).open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=COLUMNS)
            writer.writeheader()
            for row in rows:
                writer.writerow({column: row[column] for column in COLUMNS})

    def export_xlsx(
        self,
        path: Path,
        *,
        company: str | None = None,
        start_date: str | None = None,
        end_date: str | None = None,
        records: list | None = None,
    ):
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise RuntimeError("Excel dışa aktarımı için openpyxl kurulmalıdır.") from exc

        if records is not None:
            rows = records
        elif company or start_date or end_date:
            rows = self.filtered_complaints(company=company, start_date=start_date, end_date=end_date)
        else:
            rows = self.all_complaints()

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Şikâyetler"
        sheet.append(COLUMNS)

        for row in rows:
            sheet.append([row[column] for column in COLUMNS])

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(path)

    def database_stats(self) -> dict:
        """Return comprehensive metadata and statistics about the SQLite database."""
        file_size = self.db_path.stat().st_size if self.db_path.exists() else 0
        with self._connect() as conn:
            has_sentiment = self._table_exists(conn, "sentiment_results")
            total_complaints = conn.execute("SELECT COUNT(*) FROM complaints").fetchone()[0]
            total_sentiments = (
                conn.execute("SELECT COUNT(*) FROM sentiment_results").fetchone()[0]
                if has_sentiment
                else 0
            )

            company_rows = conn.execute(
                """
                SELECT 
                    company,
                    COUNT(*) as count,
                    SUM(CASE WHEN resolved = 1 THEN 1 ELSE 0 END) as resolved_count,
                    MIN(complaint_date) as min_date,
                    MAX(complaint_date) as max_date
                FROM complaints
                GROUP BY company
                ORDER BY count DESC
                """
            ).fetchall()

            companies = []
            for r in company_rows:
                comp = r["company"]
                nlp_count = 0
                if has_sentiment:
                    nlp_row = conn.execute(
                        """
                        SELECT COUNT(DISTINCT s.complaint_id)
                        FROM sentiment_results s
                        JOIN complaints c ON s.complaint_id = c.id
                        WHERE c.company = ?
                        """,
                        (comp,),
                    ).fetchone()
                    nlp_count = nlp_row[0] if nlp_row else 0

                companies.append(
                    {
                        "company": comp,
                        "count": r["count"],
                        "resolved": r["resolved_count"] or 0,
                        "nlp_count": nlp_count,
                        "min_date": r["min_date"] or "—",
                        "max_date": r["max_date"] or "—",
                    }
                )

        return {
            "db_path": str(self.db_path),
            "file_size_bytes": file_size,
            "total_complaints": total_complaints,
            "total_sentiments": total_sentiments,
            "companies": companies,
        }

    def delete_company(self, company: str) -> dict:
        """Permanently delete all complaints and associated sentiment analyses for a company."""
        norm_company = normalize_company_input(company)
        with self._connect() as conn:
            has_sentiment = self._table_exists(conn, "sentiment_results")
            has_aspect = self._table_exists(conn, "aspect_sentiment")

            ids = [
                row[0]
                for row in conn.execute(
                    "SELECT id FROM complaints WHERE company = ?", (norm_company,)
                ).fetchall()
            ]

            deleted_sentiments = 0
            deleted_aspects = 0
            if ids:
                placeholders = ",".join("?" for _ in ids)
                if has_sentiment:
                    cur = conn.execute(
                        f"DELETE FROM sentiment_results WHERE complaint_id IN ({placeholders})",
                        ids,
                    )
                    deleted_sentiments = cur.rowcount
                if has_aspect:
                    cur = conn.execute(
                        f"DELETE FROM aspect_sentiment WHERE complaint_id IN ({placeholders})",
                        ids,
                    )
                    deleted_aspects = cur.rowcount

            cur = conn.execute("DELETE FROM complaints WHERE company = ?", (norm_company,))
            deleted_complaints = cur.rowcount

        self.vacuum()
        return {
            "deleted_complaints": deleted_complaints,
            "deleted_sentiments": deleted_sentiments,
            "deleted_aspects": deleted_aspects,
        }

    def reset_database(self) -> dict:
        """Completely reset and wipe all complaints and NLP data, reclaiming disk space."""
        with self._connect() as conn:
            has_sentiment = self._table_exists(conn, "sentiment_results")
            has_aspect = self._table_exists(conn, "aspect_sentiment")
            has_sequence = self._table_exists(conn, "sqlite_sequence")

            deleted_sentiments = 0
            deleted_aspects = 0
            if has_sentiment:
                cur = conn.execute("DELETE FROM sentiment_results")
                deleted_sentiments = cur.rowcount
            if has_aspect:
                cur = conn.execute("DELETE FROM aspect_sentiment")
                deleted_aspects = cur.rowcount

            cur = conn.execute("DELETE FROM complaints")
            deleted_complaints = cur.rowcount

            if has_sequence:
                conn.execute(
                    "DELETE FROM sqlite_sequence WHERE name IN ('complaints', 'sentiment_results', 'aspect_sentiment')"
                )

        self.vacuum()
        return {
            "deleted_complaints": deleted_complaints,
            "deleted_sentiments": deleted_sentiments,
            "deleted_aspects": deleted_aspects,
        }

    def vacuum(self) -> None:
        """Reclaim unused space and defragment the SQLite database file."""
        try:
            conn = sqlite3.connect(self.db_path, isolation_level=None)
            conn.execute("VACUUM")
            conn.close()
        except sqlite3.Error:
            pass

